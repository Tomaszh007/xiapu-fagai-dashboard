# -*- coding: utf-8 -*-
"""
霞浦县发展和改革局 - 项目批复数据爬虫（v2）
数据源：霞浦县人民政府网 县发改局「政策文件」栏目
接口：POST https://www.xiapu.gov.cn/fjdzapp/search （TRS 站群检索，返回 JSON 含正文全文）

抓取目标：近半年内发布的 可研批复 / 设计批复 / 立项(项目建议书) 等新批项目
          （"同意变更/同意撤销/法人变更"类剔除，另存备查表）
字段提取：投资额 / 建设单位 / 建设地点 / 建设工期 / 资金来源 / 项目代码 / 批复正文

输出：
  霞浦发改批复_近半年.xlsx     主表（新批批复）+ 变更撤销类备查表
  approvals_recent.json        近半年明细（含正文，供 build_dashboard.py 生成看板）

运行：python crawl_xiapu.py
"""
import json
import os
import re
import time
import datetime
import requests
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ---------------- 接口常量 ----------------
BASE_URL = 'https://www.xiapu.gov.cn'
SEARCH_API = BASE_URL + '/fjdzapp/search'
CHANNEL_ID = '201588'          # 站点文档检索频道
CHNL_ID = '25042'              # 发改局 政策文件 栏目
HEADERS = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0 Safari/537.36',
    'Referer': BASE_URL + '/',
}
PAGE_SIZE = 100

# 分类规则（优先级：设计 > 可研 > 立项；命中即排除的是非项目文件）
RULES = [
    ('设计批复', [r'初步设计', r'设计概算', r'施工图设计']),
    ('可研批复', [r'可行性研究报告']),
    ('立项批复', [r'建议书', r'立项', r'项目申请报告', r'项目核准']),
]
EXCLUDE = [
    r'天然气', r'价格', r'收费', r'听证', r'普法', r'法治政府', r'观察员',
    r'意见建议', r'征求意见', r'招聘', r'考核', r'主题党日', r'党组',
    r'人事', r'预算', r'决算', r'述职', r'工作报告',
]
# 变更 / 撤销 / 法人变更类批复：对既有批复的调整，非新批项目（主表剔除，备查表保留）
CHANGE_PATTERNS = ['同意变更', '同意撤销', '法人变更']
# 公文尾部噪声（印章/落款区）截断
TAIL_RE = re.compile(r'(附件[:：]|霞浦县发展和改革局\s*[\s　]*\n?\s*20\d\d年|（此件主动公开）)')

# 时间范围（近半年）
TODAY = datetime.date.today()
DATE_FROM = (TODAY - datetime.timedelta(days=183)).isoformat()

# 输出目录（脚本所在目录）
HERE = os.path.dirname(os.path.abspath(__file__))
OUT_XLSX = os.path.join(HERE, '霞浦发改批复_近半年.xlsx')
OUT_JSON = os.path.join(HERE, 'approvals_recent.json')


# ---------------- 抓取 ----------------
def fetch_all():
    """POST 翻页抓取栏目全部文档（含正文），按 docid/url 去重"""
    out, seen, page = [], set(), 1
    while True:
        for attempt in range(3):
            try:
                r = requests.post(SEARCH_API, data={
                    'channelid': CHANNEL_ID,
                    'sortfield': '-docorderpri,-docreltime',
                    'classsql': f'chnlid={CHNL_ID}',
                    'classnum': 100, 'classsort': 0,
                    'cache': 'true', 'page': page, 'prepage': PAGE_SIZE,
                }, headers=HEADERS, timeout=45)
                r.raise_for_status()
                d = r.json()
                break
            except Exception as e:
                if attempt == 2:
                    raise RuntimeError(f'第{page}页请求失败: {e}')
                time.sleep(2 * (attempt + 1))
        if d.get('error'):
            raise RuntimeError(f'接口返回错误 page={page}: {d.get("error")}')
        docs = d.get('data') or []
        out.extend(docs)
        pages = int(d.get('pageCount') or 1)
        if page >= pages or not docs:
            break
        page += 1
        time.sleep(0.6)
    # 去重
    uniq, useen = [], set()
    for doc in out:
        key = str(doc.get('docid2') or doc.get('docpuburl'))
        if key in useen:
            continue
        useen.add(key)
        uniq.append(doc)
    return uniq


# ---------------- 清洗与提取 ----------------
def clean_body(txt, limit=None):
    if not txt:
        return ''
    t = txt.replace('\r', '')
    t = re.sub(r'[ \t\u3000]+', ' ', t)
    t = re.sub(r'\n{3,}', '\n\n', t).strip()
    m = TAIL_RE.search(t)
    if m and m.start() > 120:
        t = t[:m.start()].strip()
    if limit:
        t = t[:limit]
    return t


def pick(pattern, text, group=1):
    m = re.search(pattern, text)
    return m.group(group).strip(' 　:：。') if m else ''


def classify(title):
    """返回 (类别, 是否项目批复)"""
    for p in EXCLUDE:
        if re.search(p, title or ''):
            return None, False
    for cat, pats in RULES:
        if any(re.search(p, title or '') for p in pats):
            return cat, True
    return None, False


def is_change_approval(title):
    return any(p in (title or '') for p in CHANGE_PATTERNS)


def normalize_date(s):
    s = (s or '').replace('.', '-')
    m = re.search(r'(\d{4})年(\d{1,2})月(\d{1,2})日', s)
    if m:
        return f'{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}'
    return s[:10]


def extract(rec, cat):
    """从标题 + 正文抽取结构化字段"""
    title = (rec.get('doctitle') or '').strip()
    body = clean_body(rec.get('doccontent'), limit=6000)

    # 项目名称：标题取主体，去掉机构前缀与文种后缀
    proj = ''
    m = re.search(r'关于(.+?)(?:的批复|批复|$)', title)
    if m:
        proj = m.group(1)
    if not proj:
        proj = pick(r'项目名称[：:]\s*\n?\s*(.+)', body)
    proj = re.sub(r'^(?:同意)?(?:变更|撤销|调整|废止|终止)', '', proj)
    for p in ('建议书暨可行性研究报告', '建议书及可行性研究报告',
              '可行性研究报告', '项目建议书', '初步设计及概算', '初步设计',
              '项目申请报告', '投资估算'):
        proj = proj.replace(p, '')
    proj = re.sub(r'(调整|变更)$', '', proj)
    proj = proj.strip('的 。、暨')
    if not proj:
        proj = title

    # 项目代码（福建投资项目在线审批监管平台编码）
    pcode = pick(r'项目(?:代码|编码)[：:]\s*([0-9A-Za-z\-]{10,})', body)
    if not pcode:
        pcode = pick(r'([2]\d{3}-\d{6}-\d{2}-\d{2}-\d{6})', title + body)

    # 建设单位
    unit = pick(r'(?:项目|建设)单位[：:]\s*\n?\s*([^\n。；]{2,40})', body)
    if not unit:
        m2 = re.match(r'([^\s：:，,]{2,30}(?:政府|局|公司|委员会|办事处|中心|医院|学校|集团))[：:]', body)
        unit = m2.group(1) if m2 else ''

    # 投资额（万元）：总投资匡算（建议书）> 概算投资（调整批复）> 总投资 > 投资估算；亿元自动换算
    invest_wan = ''
    for pat, is_yi in [
        (r'总投资匡算(?:及资金来源)?[：:]?\s*([0-9]+(?:\.[0-9]+)?)\s*亿元', True),
        (r'总投资匡算(?:及资金来源)?[：:]?\s*([0-9][0-9,\.]*)\s*万元', False),
        (r'概算投资[：:]?\s*([0-9]+(?:\.[0-9]+)?)\s*亿元', True),
        (r'概算投资[：:]?\s*([0-9][0-9,\.]*)\s*万元', False),
        (r'总投资(?:为|约|共计|估算为)?\s*([0-9]+(?:\.[0-9]+)?)\s*亿元', True),
        (r'总投资(?:为|约|共计|估算为)?\s*([0-9][0-9,\.]*)\s*万元', False),
        (r'投资估算(?:为|约|共计)?\s*([0-9]+(?:\.[0-9]+)?)\s*亿元', True),
        (r'投资估算(?:为|约|共计)?\s*([0-9][0-9,\.]*)\s*万元', False),
    ]:
        v = pick(pat, body)
        if v:
            v = v.rstrip('.').replace(',', '')
            try:
                invest_wan = f'{float(v) * 10000:.2f}'.rstrip('0').rstrip('.') if is_yi else v
                break
            except ValueError:
                invest_wan = ''

    location = pick(r'(?:项目)?建设(?:地点|地址)[：:]\s*\n?\s*([^\n。]{2,40})', body)
    duration = pick(r'(?:项目建设)?(?:工期|建设工期)[：:]?\s*(?:为)?\s*([0-9]+\s*(?:个)?月)', body)
    # 资金来源：先取"资金来源为XXX"（建议书句式，止于逗号/括号），再取"资金来源：XXX"
    fund = pick(r'资金来源为\s*([^，。；（\n]{2,60})', body)
    if not fund:
        fund = pick(r'资金来源[：:]\s*([^\n。；（]{2,80})', body)
    if not fund:
        fund = pick(r'资金(?:由|系)\s*([^\n。]{2,60})', body)
    # 误捕投资额句时置空
    if fund and re.search(r'万元|亿元|总投资|匡算|估算', fund):
        fund = ''

    return {
        'category': cat,
        'project_name': proj,
        'project_code': pcode,
        'build_unit': unit,
        'location': location,
        'invest_wan': invest_wan,
        'duration': duration,
        'fund_source': fund,
        'title': title,
        'doc_no': (rec.get('fileno') or '').strip(),
        'publish_date': (rec.get('docreltime') or '').replace('.', '-')[:10],
        'doc_date': normalize_date(rec.get('pubdate')),
        'url': rec.get('docpuburl') or rec.get('chnldocurl') or '',
        'doc_id': str(rec.get('docid2') or ''),
        'content': body,
    }


def to_num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


# ---------------- 输出 ----------------
THIN = Border(*[Side(style='thin', color='B0B0B0')] * 4)
HEAD_FONT = Font(name='微软雅黑', bold=True, color='FFFFFF', size=11)
HEAD_FILL = PatternFill('solid', fgColor='2F5597')


def _fill_sheet(ws, rows, title):
    ws.title = title
    headers = ['项目名称', '批复类型', '文号', '投资额(万元)', '建设单位', '建设地点', '建设工期', '资金来源', '成文日期', '发布日期', '原文链接']
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center')
    for r in rows:
        ws.append([
            r['project_name'], r['category'], r['doc_no'],
            to_num(r['invest_wan']) or '',
            r['build_unit'], r['location'], r['duration'], r['fund_source'],
            r['doc_date'], r['publish_date'], r['url'],
        ])
    for i in range(2, len(rows) + 2):
        for c in range(1, len(headers) + 1):
            ws.cell(row=i, column=c).alignment = Alignment(vertical='top', wrap_text=(c == 1))
            ws.cell(row=i, column=c).border = THIN
    widths = [46, 12, 18, 13, 22, 14, 10, 22, 12, 12, 50]
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.freeze_panes = 'A2'


def write_excel(rows, excluded):
    wb = Workbook()
    _fill_sheet(wb.active, rows, '批复项目')
    _fill_sheet(wb.create_sheet(), excluded, '变更撤销类备查')
    wb.save(OUT_XLSX)


def main():
    docs = fetch_all()
    print(f'栏目共抓取 {len(docs)} 篇（含全文）')

    # 近半年发布 + 分类 + 变更类分流
    rows, excluded, stat = [], [], {}
    for doc in docs:
        pub = (doc.get('docreltime') or '').replace('.', '-')[:10]
        if not pub or pub < DATE_FROM:
            continue
        title = doc.get('doctitle') or ''
        # 变更 / 撤销 / 法人变更类：先行分流备查表（撤销/法人变更不命中分类规则）
        if is_change_approval(title):
            cat = classify(title)[0] or '变更撤销类'
            excluded.append(extract(doc, cat))
            continue
        cat, ok = classify(title)
        if not ok:
            continue
        rows.append(extract(doc, cat))
        stat[cat] = stat.get(cat, 0) + 1

    # 排序：发布日期倒序
    rows.sort(key=lambda r: r.get('publish_date', ''), reverse=True)
    excluded.sort(key=lambda r: r.get('publish_date', ''), reverse=True)

    write_excel(rows, excluded)
    with open(OUT_JSON, 'w', encoding='utf-8') as f:
        json.dump({'date_from': DATE_FROM, 'date_to': TODAY.isoformat(),
                   'items': rows, 'excluded': excluded}, f, ensure_ascii=False, indent=1)

    invest_items = [r for r in rows if to_num(r['invest_wan']) > 0]
    total_inv = sum(to_num(r['invest_wan']) for r in rows)
    print(f'近半年新批批复 {len(rows)} 条（可研{stat.get("可研批复",0)}/设计{stat.get("设计批复",0)}/立项{stat.get("立项批复",0)}），'
          f'剔除变更/撤销/法人变更类 {len(excluded)} 条')
    print(f'投资额提取 {len(invest_items)}/{len(rows)} 条，合计约 {total_inv/10000:.2f} 亿元')
    print(f'Excel: {OUT_XLSX}')
    print(f'JSON : {OUT_JSON}')


if __name__ == '__main__':
    main()